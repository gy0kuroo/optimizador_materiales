from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from cutless.exports.excel import generar_excel
from cutless.exports.pdf import generar_pdf
from cutless.models import Optimizacion
from cutless.render import generar_grafico
from cutless.services.optimization import convertir_info_desperdicio_unidad
from cutless.units import obtener_factor_area_desde_cm2


class FactoresAreaTests(SimpleTestCase):
    def test_equivalencias_de_un_metro_cuadrado(self):
        esperados = {'cm': 10000, 'm': 1, 'mm': 1000000, 'in': 1550.0031000062, 'ft': 10.76391041671}
        for unidad, esperado in esperados.items():
            with self.subTest(unidad=unidad):
                self.assertAlmostEqual(10000 * obtener_factor_area_desde_cm2(unidad), esperado, places=8)

    def test_alias_y_unidad_desconocida(self):
        self.assertEqual(obtener_factor_area_desde_cm2('pulgadas'), obtener_factor_area_desde_cm2('in'))
        self.assertEqual(obtener_factor_area_desde_cm2('desconocida'), 1)

    def test_totales_y_tableros_se_convierten_sin_alterar_original(self):
        info = {
            'area_usada_total': 10000, 'desperdicio_total': 10000,
            'info_tableros': [{'numero': 1, 'area_usada': 10000, 'desperdicio': 10000, 'num_piezas': 2}],
            'num_tableros': 1, 'num_piezas_colocadas': 2,
        }
        for unidad, esperado in [('ft', 10.76), ('in', 1550.00), ('cm', 10000), ('m', 1), ('mm', 1000000)]:
            with self.subTest(unidad=unidad):
                convertido = convertir_info_desperdicio_unidad(info, unidad)
                self.assertEqual(convertido['area_usada_total'], esperado)
                self.assertEqual(convertido['desperdicio_total'], esperado)
                self.assertEqual(convertido['info_tableros'][0]['area_usada'], esperado)
                self.assertEqual(convertido['info_tableros'][0]['desperdicio'], esperado)
                self.assertEqual(convertido['num_piezas_colocadas'], 2)
        self.assertEqual(info['area_usada_total'], 10000)
        self.assertEqual(info['info_tableros'][0]['area_usada'], 10000)


class ExportacionesAreaTests(TestCase):
    def setUp(self):
        usuario = User.objects.create_user('areas_user')
        self.opt = Optimizacion.objects.create(
            usuario=usuario, ancho_tablero=100, alto_tablero=100,
            unidad_medida='ft', piezas='Cuadrado,1,1,2', margen_corte=0,
        )
        self.info = {
            'area_usada_total': 1858.0608, 'desperdicio_total': 8141.9392,
            'info_tableros': [{'numero': 1, 'area_usada': 1858.0608, 'desperdicio': 8141.9392, 'porcentaje_uso': 18.58, 'num_piezas': 2}],
            'num_tableros': 1, 'num_piezas_solicitadas': 2, 'num_piezas_colocadas': 2,
        }

    def test_excel_areas_unitarias_y_totales(self):
        for unidad in ('cm', 'm', 'mm', 'in', 'ft'):
            with self.subTest(unidad=unidad):
                self.opt.unidad_medida = unidad
                wb = load_workbook(generar_excel(self.opt, {}, [{'nombre': 'Cuadrado', 'ancho': 1, 'alto': 1, 'cantidad': 2}]))
                ws = wb['Piezas']
                self.assertEqual(ws['E2'].value, f'1.0 {unidad}²')
                self.assertEqual(ws['F2'].value, f'2.0 {unidad}²')
                self.assertEqual(ws['F3'].value, f'2.0 {unidad}²')
                wb.close()

    def test_grafico_muestra_area_correcta_en_pies(self):
        from matplotlib.axes import Axes
        original = Axes.text
        textos = []
        def capturar(ax, x, y, texto, *args, **kwargs):
            textos.append(str(texto))
            return original(ax, x, y, texto, *args, **kwargs)
        with patch.object(Axes, 'text', autospec=True, side_effect=capturar):
            imagenes, _, info = generar_grafico([(30.48, 30.48, 2)], 100, 100, unidad='ft', margen_corte=0)
        self.assertEqual(len(imagenes), 1)
        self.assertIn('Área usada: 2.0 ft²', '\n'.join(textos))
        self.assertAlmostEqual(info['area_usada_total'], 1858.0608)

    def test_pdf_muestra_areas_unitarias_totales_y_tableros(self):
        from reportlab.pdfgen.canvas import Canvas
        original = Canvas.drawString
        textos = []
        def capturar(canvas, x, y, texto, *args, **kwargs):
            textos.append(str(texto))
            return original(canvas, x, y, texto, *args, **kwargs)
        imagenes, _, _ = generar_grafico([(30.48, 30.48, 2)], 100, 100, margen_corte=0)
        with patch.object(Canvas, 'drawString', autospec=True, side_effect=capturar):
            archivo = generar_pdf(self.opt, imagenes, info_desperdicio=self.info)
        self.assertTrue((Path(settings.MEDIA_ROOT) / archivo).is_file())
        self.assertIn('1.0 ft²', textos)
        self.assertIn('2.0 ft²', textos)
        self.assertIn('Total: 2.0 ft²', textos)
        self.assertGreaterEqual(textos.count('2.0 ft²'), 2)
        self.assertIn('8.76 ft²', textos)
